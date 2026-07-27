"""routing_lane — 가로 실행 레인 + 도중 판정에서 빠지는 분기 (골든 s06에서 꺼냄, 2026-07-26).

출처: `s06_variants.variant_c`(골든 S6)의 실행 그래프를 부품 계약 3조건으로 옮긴 것.
좌표·도형·색은 골든 실측 그대로여서 `compare_golden` 픽셀 동일로 무손실을 증명한다.

## 이 도해가 무엇을 말하나

입력이 **한 줄로 흐르는데**, 도중의 판정 하나에서 일부가 **본선 밖으로 빠진다**. 통과 경로는
가로로 끝까지 이어지고 거절 경로는 아래로 떨어져 거기서 끝난다 — 방향이 다르다는 것 자체가
"본선에 진입하지 못했다"를 말한다. 단계를 세로로 쌓거나 거절을 옆 칸에 두면 그 뜻이 사라진다.

## 노드 종류가 폭과 모양을 정한다

    start    좁은 둥근 사각 · 테두리만      입력 단자
    process  넓은 둥근 사각 · 옅은 채움     실제 일을 하는 단계
    gate     마름모 · accent 테두리 · 더 큼  갈리는 곳(레인에서 유일하게 튀어나온다)
    end      좁은 둥근 사각 · 진한 채움     도달점

폭이 종류에서 나오므로 노드 수는 자유다. 간격은 남는 자리를 노드 수로 나눠 파생한다 —
좌표 리터럴이 없어서 6노드든 9노드든 같은 코드가 그린다.
"""

from pptx.enum.text import PP_ALIGN

from ..kit import add_box, add_text, set_shape_text
from . import elements as E
from . import Box, merged, require

META = {
    "name": "routing_lane",
    "어휘": "G04",  # archetype-catalog 등재 어휘(routing_band)와 대응
    "물성": "한 줄로 흐르다 도중 판정에서 일부가 본선 밖으로 빠진다(통과 경로 vs 거절 경로)",
    "arity": "파생 — 노드 폭은 종류에서, 간격은 남는 자리에서 나온다",
    # 셀 수 있는 선택 조건 — 컴포저는 재료를 세어 이것과 **비교만** 한다(산문 해석 없음).
    "accepts": {'sets': 1, 'flow': '순차', 'order': True, 'extra': True, 'n': (3, 9), 'ends': False},
    "keys": ["nodes"],
    "optional": ["branch", "in_label"],
}

LAYOUT = {
    # box 원점 기준 offset(inch). 골든 s06 실측값.
    "node_h": 0.60,
    "gate_extra": 0.25,  # 판정만 위아래로 더 크다 — 레인에서 튀어나와야 "여기서 갈린다"가 보인다
    "w_start": 1.00,
    "w_process": 1.45,
    "w_gate": 1.15,
    "w_end": 1.00,
    "tag_dy": 0.825,  # box top → 노드 아래 태그 top
    "tag_h": 0.24,
    "tag_bleed": 0.35,  # 태그는 노드보다 좌우로 넓다(긴 태그가 줄바꿈되지 않게)
    "in_gap": 0.03,  # 판정 우변 → IN 라벨
    "in_dy": -0.32,  # 레인 중앙 기준
    "label_w": 0.50,
    "label_h": 0.22,
    "branch_dy": 1.325,  # box top → 거절 박스 top
    "branch_w": 2.00,
    "branch_h": 0.45,
    "out_gap": 0.12,  # 판정 하단 → OUT 라벨
    "desc_w": 4.20,  # 거절 설명 폭 — 자리 전체가 아니다(길면 레인 밖까지 흘러 읽기가 끊긴다)
    "desc_h": 0.26,
}

_KIND_W = {"start": "w_start", "process": "w_process", "gate": "w_gate", "end": "w_end"}


def _widths(nodes, L):
    return [L[_KIND_W[n.get("kind", "process")]] for n in nodes]


def measure(data, kit, layout=None):  # noqa: ARG001 — 계약 시그니처 고정
    """필요 폭·높이(inch). 노드가 붙어버리지 않는 최소 폭을 낸다."""
    L = merged(LAYOUT, layout)
    ws = _widths(data["nodes"], L)
    return sum(ws) + (len(ws) - 1) * 0.20, L["branch_dy"] + L["branch_h"]


def draw(slide, box, data, kit, layout=None):
    require(data, META["keys"], META["name"])
    box = Box(*box)  # 4튜플도 받는다 — 목업 러너·장이 둘 다 쓴다
    L = merged(LAYOUT, layout)
    C, S, F = kit["rgb"], kit["sizes"], kit["fonts"]
    nodes = data["nodes"]

    kinds = [n.get("kind", "process") for n in nodes]
    bad = sorted(set(kinds) - set(_KIND_W))
    if bad:
        raise ValueError(
            f"routing_lane: 모르는 노드 종류 {bad} — {sorted(_KIND_W)} 중 하나여야 한다"
        )
    if data.get("branch") and kinds.count("gate") != 1:
        raise ValueError(
            f"routing_lane: 분기를 그리려면 gate가 정확히 1개여야 하는데 {kinds.count('gate')}개다 "
            "— 어느 판정에서 빠지는지가 정해지지 않으면 도해가 거짓말을 한다"
        )

    ws = _widths(nodes, L)
    if len(nodes) < 2:
        raise ValueError("routing_lane: 노드가 2개 미만이면 레인이 아니다")
    gap = (box.w - sum(ws)) / (len(nodes) - 1)
    if gap < 0.12:
        raise ValueError(
            f'routing_lane: 노드 {len(nodes)}개가 자리 {box.w:.2f}"에 안 들어간다(간격 {gap:.2f}") '
            "— 폭을 줄이면 글자가 넘치므로 노드를 줄이거나 자리를 넓힌다"
        )

    xs, x = [], box.x
    for w in ws:
        xs.append(x)
        x += w + gap
    node_h = L["node_h"]
    mid = box.y + (node_h + L["gate_extra"]) / 2

    # ── 노드 ──
    for kind, node, nx, w in zip(kinds, nodes, xs, ws, strict=True):
        gate = kind == "gate"
        h = node_h + L["gate_extra"] if gate else node_h
        fill, line, ink = {
            "start": (C["bg"], C["muted"], C["primary"]),
            "process": (C["bg_alt"], None, C["primary"]),
            "gate": (C["bg"], C["accent"], C["primary"]),
            "end": (C["primary"], None, C["bg"]),
        }[kind]
        shp = add_box(
            slide,
            nx,
            mid - h / 2,
            w,
            h,
            fill=fill,
            line=line,
            line_w=1.25 if line else None,
            shape="diamond" if gate else "round",
        )
        set_shape_text(
            shp, node["name"], S["caption" if gate else "body"], F["head"], ink, bold=True
        )

    for i in range(len(nodes) - 1):
        E.arrow(slide, kit, xs[i] + ws[i], mid, xs[i + 1], mid)

    # ── 노드 아래 태그 (단계가 실제로 무엇을 하나) ──
    for node, nx, w in zip(nodes, xs, ws, strict=True):
        if not node.get("tag"):
            continue
        bleed = L["tag_bleed"]
        add_text(
            slide,
            nx - bleed,
            box.y + L["tag_dy"],
            w + bleed * 2,
            L["tag_h"],
            node["tag"],
            S["caption"],
            F["head"],
            C["primary"],
            bold=True,
            align=PP_ALIGN.CENTER,
        )

    if "gate" not in kinds:
        return
    gi = kinds.index("gate")
    gate_cx = xs[gi] + ws[gi] / 2
    gate_bot = mid + (node_h + L["gate_extra"]) / 2

    if data.get("in_label"):
        add_text(
            slide,
            xs[gi] + ws[gi] + L["in_gap"],
            mid + L["in_dy"],
            L["label_w"],
            L["label_h"],
            data["in_label"],
            S["caption"],
            F["head"],
            C["accent"],
            bold=True,
        )

    # ── 분기 — 본선에서 아래로 떨어져 거기서 끝난다 ──
    br = data.get("branch")
    if not br:
        return
    bw, bh = L["branch_w"], L["branch_h"]
    bx, by = gate_cx - bw / 2, box.y + L["branch_dy"]
    shp = add_box(slide, bx, by, bw, bh, fill=C["bg_alt"], shape="round")
    set_shape_text(shp, br["box"], S["caption"], F["head"], C["primary"], bold=True)
    E.arrow(slide, kit, gate_cx, gate_bot, gate_cx, by)
    if br.get("label"):
        add_text(
            slide,
            gate_cx + L["out_gap"],
            gate_bot + L["out_gap"],
            L["label_w"] + 0.10,
            L["label_h"],
            br["label"],
            S["caption"],
            F["head"],
            C["accent"],
            bold=True,
        )
    if br.get("desc"):
        add_text(
            slide,
            bx + bw + 0.2,
            by + 0.1,
            L["desc_w"],
            L["desc_h"],
            br["desc"],
            S["caption"],
            F["body"],
            C["muted"],
        )


SAMPLES = [
    (
        "최소 3노드",
        {
            "nodes": [
                {"name": "질문", "kind": "start"},
                {"name": "판정", "kind": "gate"},
                {"name": "응답", "kind": "end"},
            ],
            "branch": {"box": "거절 메시지", "label": "OUT"},
            "in_label": "IN",
        },
    ),
    (
        "골든 7노드",
        {
            "nodes": [
                {"name": "질문", "kind": "start"},
                {"name": "Analyze", "tag": "용어사전 매칭"},
                {"name": "판정", "kind": "gate"},
                {"name": "Retrieve", "tag": "그래프 1홉 탐색"},
                {"name": "Generate", "tag": "판단트리 주입"},
                {"name": "Format", "tag": "경고·꼬리질문"},
                {"name": "응답", "kind": "end"},
            ],
            "branch": {
                "box": "거절 메시지",
                "label": "OUT",
                "desc": "범위 밖 질문은 본선에 진입하지 못하고 즉시 거절된다",
            },
            "in_label": "IN",
        },
    ),
]
